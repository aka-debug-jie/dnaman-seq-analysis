"""Build an Excel workbook summarising everything under results/."""

import os
import re

from . import config

METRIC_PATTERNS = [
    ("length_bp", re.compile(r"SEQ\s+\S+:\s*(\d+)\s*bp")),
    ("composition", re.compile(r"Composition\s+(\d+)\s+A;\s*(\d+)\s+C;\s*(\d+)\s+G;\s*(\d+)\s+T")),
    ("sites_found", re.compile(r"Screened with \d+ enzymes,\s*(\d+) sites found")),
    ("enzymes_cut", re.compile(r"^([A-Za-z]+)\s+\d+\s+\S+/\S+", re.M)),
    ("longest_orf_aa", re.compile(r"Plus\s+\d+\s+(\d+)\s+(\d+)-(\d+)")),
    ("pcr_pairs", re.compile(r"^\s*\d+\s+[ACGT]+\s+[\d.]+\S*\s+and\s+\d+\s+[ACGT]+", re.M)),
    ("consensus_len", re.compile(r"assembly_consensus len=(\d+)")),
    ("identity", re.compile(r"identity=\s*(\d+)%")),
    ("aa_length", re.compile(r"Protein Length=(\d+)")),
    ("pi", re.compile(r"Predicted pI=([\d.]+)")),
]


def _extract(text):
    out = {}
    for name, rx in METRIC_PATTERNS:
        m = rx.search(text)
        if not m:
            continue
        if name == "composition":
            out["A"] = int(m.group(1))
            out["C"] = int(m.group(2))
            out["G"] = int(m.group(3))
            out["T"] = int(m.group(4))
        elif name == "enzymes_cut":
            enz = sorted(set(rx.findall(text)))
            out["enzymes_cut"] = ", ".join(enz)
        elif name == "longest_orf_aa":
            out["longest_orf_aa"] = int(m.group(1))
            out["longest_orf_pos"] = "%s-%s" % (m.group(2), m.group(3))
        elif name == "pcr_pairs":
            out["pcr_pairs"] = len(rx.findall(text))
        elif name == "identity":
            out["identity_pct"] = int(m.group(1))
        elif name in ("aa_length", "consensus_len", "sites_found", "length_bp"):
            out[name] = int(m.group(1))
        elif name == "pi":
            out["pI"] = float(m.group(1))
    return out


def collect(results_dir=None):
    """-> list of dicts: topic, file, kind, bytes, metrics"""
    results_dir = results_dir or config.RESULTS
    rows = []
    for root, dirs, files in os.walk(results_dir):
        dirs[:] = [d for d in dirs if not d.startswith("99_")]
        if root == results_dir:
            continue
        topic = os.path.relpath(root, results_dir).replace("\\", "/")
        if topic.startswith("99_"):
            continue
        for fname in sorted(files):
            fpath = os.path.join(root, fname)
            ext = os.path.splitext(fname)[1].lower()
            metrics = {}
            if ext == ".txt":
                try:
                    with open(fpath, encoding="utf-8", errors="replace") as fh:
                        metrics = _extract(fh.read())
                except Exception:
                    pass
            rows.append({
                "topic": topic,
                "file": fname,
                "kind": "figure" if ext == ".png" else "text",
                "bytes": os.path.getsize(fpath),
                "metrics": metrics,
            })
    return rows


def build_excel(out_path=None, results_dir=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    rows = collect(results_dir)
    out_path = out_path or os.path.join(config.RESULTS, "summary.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "files"
    ws.append(["topic", "file", "kind", "bytes"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r["topic"], r["file"], r["kind"], r["bytes"]])
    ws.freeze_panes = "A2"
    for col, width in zip("ABCD", (26, 44, 10, 12)):
        ws.column_dimensions[col].width = width

    ws2 = wb.create_sheet("metrics")
    keys = ["length_bp", "A", "C", "G", "T", "sites_found", "enzymes_cut",
            "longest_orf_aa", "longest_orf_pos", "pcr_pairs", "identity_pct",
            "aa_length", "pI", "consensus_len"]
    ws2.append(["topic", "file"] + keys)
    for c in ws2[1]:
        c.font = Font(bold=True)
    for r in rows:
        if not r["metrics"]:
            continue
        ws2.append([r["topic"], r["file"]] + [r["metrics"].get(k) for k in keys])
    ws2.freeze_panes = "A2"
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 40
    for i in range(len(keys)):
        ws2.column_dimensions[chr(ord("C") + i)].width = 13

    ws3 = wb.create_sheet("summary")
    ws3.append(["topic", "files", "texts", "figures", "total_kb"])
    for c in ws3[1]:
        c.font = Font(bold=True)
    topics = {}
    for r in rows:
        t = topics.setdefault(r["topic"], {"files": 0, "text": 0, "figure": 0, "bytes": 0})
        t["files"] += 1
        t[r["kind"]] += 1
        t["bytes"] += r["bytes"]
    for topic, t in sorted(topics.items()):
        ws3.append([topic, t["files"], t["text"], t["figure"], round(t["bytes"] / 1024, 1)])
    for col, width in zip("ABCDE", (28, 8, 8, 9, 10)):
        ws3.column_dimensions[col].width = width

    out_dir = os.path.dirname(out_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(out_path)
    return out_path, len(rows), len(topics)


def main(argv=None):
    import argparse

    p = argparse.ArgumentParser(prog="dnaman-report", description="summarise results into xlsx")
    p.add_argument("--out", default=None)
    p.add_argument("--results", default=None)
    a = p.parse_args(argv)
    path, files, topics = build_excel(a.out, a.results)
    print("workbook : %s" % path)
    print("rows     : %d file(s) across %d topic(s)" % (files, topics))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
